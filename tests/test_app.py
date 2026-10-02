from pathlib import Path
from io import BytesIO
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from thermo import *
from chart import svg_chart, png_chart
from exports import excel_export

ROOT=Path(__file__).parents[1]


def button(app, label):
    return next(b for b in app.button if b.label==label)


def test_app_examples_reset_and_exports():
    app=AppTest.from_file(str(ROOT/'app.py'), default_timeout=45).run()
    assert not app.exception
    button(app,'Beispiel Sommer').click().run()
    assert not app.exception
    assert len(app.session_state.rows)==3
    assert app.session_state.rows[1]['process']==COOL
    button(app,'Bild und Excel vorbereiten').click().run()
    assert not app.exception
    png,excel,svg=app.session_state.export_bundle
    assert png[:8]==b'\x89PNG\r\n\x1a\n'
    ET.fromstring(svg)
    wb=load_workbook(BytesIO(excel))
    assert wb.sheetnames==['Zustände','Prozesse','Modell und Hinweise']
    assert wb['Zustände'].max_row==4
    button(app,'Beispiel Winter').click().run()
    assert not app.exception
    assert app.session_state.rows[0]['values']==[-10.,80.]
    assert app.session_state.export_bundle is None
    app.number_input(key='pressure').set_value(1013.25).run()
    button(app,'↺ 950 hPa').click().run()
    assert app.session_state.pressure==950
    button(app,'Punkte leeren').click().run()
    assert app.session_state.rows==[]


def test_app_add_both_modes_and_reject_supersaturation():
    app=AppTest.from_file(str(ROOT/'app.py'), default_timeout=45).run()
    button(app,'＋ Zustand hinzufügen').click().run()
    assert len(app.session_state.rows)==1
    next(r for r in app.radio if r.label=='Eingabeweg').set_value('Prozess vorgeben').run()
    next(s for s in app.selectbox if s.label=='Zustandsänderung').set_value(COOL).run()
    next(n for n in app.number_input if n.label=='Zieltemperatur [°C]').set_value(5.).run()
    button(app,'＋ Prozess und Zielpunkt hinzufügen').click().run()
    assert not app.exception
    assert len(app.session_state.rows)==2
    assert app.session_state.rows[-1]['process']==COOL
    next(n for n in app.number_input if n.label=='Zieltemperatur [°C]').set_value(-10.).run()
    assert button(app,'＋ Prozess und Zielpunkt hinzufügen').disabled


def test_excel_user_names_are_not_formulas():
    s=solve(('T','phi'),(20,50))
    rows=[{'name':'=1+1','pair':['T','phi'],'values':[20,50],'process':FREE}]
    wb=load_workbook(BytesIO(excel_export(rows,[s],950)))
    assert wb['Zustände']['B2'].data_type=='s'
    assert wb['Zustände']['B2'].value=='=1+1'


def test_table_reorder_delete_and_invalid_edit_are_atomic():
    app=AppTest.from_file(str(ROOT/'app.py'), default_timeout=45).run()
    button(app,'Beispiel Sommer').click().run()
    key=f'points_{app.session_state.revision}'
    app.session_state[key]={'edited_rows':{0:{'Reihenfolge':2},1:{'Reihenfolge':1},2:{'Löschen':True}},'added_rows':[],'deleted_rows':[]}
    button(app,'Tabellenänderungen übernehmen').click().run()
    assert not app.exception
    assert [r['name'] for r in app.session_state.rows]==['Nach Kühler','Außenluft']
    key=f'points_{app.session_state.revision}'
    original=app.session_state.rows[0]['values'].copy()
    app.session_state[key]={'edited_rows':{0:{'Wert 1':99.}},'added_rows':[],'deleted_rows':[]}
    button(app,'Tabellenänderungen übernehmen').click().run()
    assert app.session_state.rows[0]['values']==original
    assert len(app.error)>0


def test_pressure_change_invalidates_process_and_explicit_recalc_repairs():
    app=AppTest.from_file(str(ROOT/'app.py'), default_timeout=45).run()
    button(app,'Beispiel Sommer').click().run()
    app.number_input(key='pressure').set_value(1013.25).run()
    assert not app.exception
    assert button(app,'Bild und Excel vorbereiten').disabled
    next(b for b in app.button if b.key=='recalc_1').click().run()
    next(b for b in app.button if b.key=='recalc_2').click().run()
    assert not app.exception
    assert not button(app,'Bild und Excel vorbereiten').disabled
